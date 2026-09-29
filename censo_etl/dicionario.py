"""Leitura do Dicionário de Variáveis do INEP (.xlsx).

Cada aba reconhecida vira uma `TabelaDicionario`. O layout esperado de uma aba:

    ... linhas de título ("Dicionário de Variáveis - Tabela de Escola")
    N | Nome da Variável | Descrição | Tipo | Tamanho | Categoria | Coleta por ano ... | Notas
      |                  |           |      |         |           | 19 | 20 | ... | 25 |
    1 | NU_ANO_CENSO     | ...

A matriz "Coleta por ano" começa na coluna 6 e vai até a primeira célula vazia
da linha seguinte ao cabeçalho; a coluna seguinte é a de notas.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from censo_etl.nomes import identificar_tabela
from censo_etl.tabelas import CHAVES
from censo_etl.texto import espacos_normalizados

log = logging.getLogger(__name__)

_CODIGO_RE = re.compile(r"^\s*(-?\d+)\s*-\s*(\S.*)$")         # "1 - Federal"
_NOTA_CATEGORIA_RE = re.compile(r"^\s*-\s+(\S.*)$")          # "- Não se aplica a ..."
_COLUNA_PRIMEIRO_ANO = 6
TITULO_PADRAO = "Dicionário de Variáveis"


@dataclass
class Variavel:
    ordem: int
    nome: str
    descricao: str
    tipo: str                               # "Num", "Char" ou "Data"
    tamanho: int | float | None             # como veio da planilha
    rotulos_valor: dict = field(default_factory=dict)   # código → rótulo
    notas_aplicabilidade: list[str] = field(default_factory=list)
    anos_coleta: dict[str, str | None] = field(default_factory=dict)  # "2025" → "s"/"n"
    notas: str | None = None


@dataclass
class TabelaDicionario:
    chave: str
    titulo: str
    variaveis: list[Variavel]
    ano_dicionario: str | None   # última coluna de "Coleta por ano"


def parse_categoria(texto) -> tuple[dict[int, str], list[str]]:
    """Separa a célula de categorias em `{código: rótulo}` e notas de aplicabilidade.

    Linha sem padrão é continuação da anterior.
    """
    entradas: list[list] = []   # [tipo, código, texto]
    for linha in str(texto or "").split("\n"):
        linha = linha.rstrip()
        if not linha.strip():
            continue
        if m := _CODIGO_RE.match(linha):
            entradas.append(["codigo", int(m.group(1)), m.group(2).strip()])
        elif m := _NOTA_CATEGORIA_RE.match(linha):
            entradas.append(["nota", None, m.group(1).strip()])
        elif entradas:
            entradas[-1][2] = (entradas[-1][2] + " " + linha.strip()).strip()
        else:
            entradas.append(["nota", None, linha.strip()])
    rotulos = {codigo: txt for tipo, codigo, txt in entradas if tipo == "codigo"}
    notas = [txt for tipo, _, txt in entradas if tipo == "nota"]
    return rotulos, notas


def _indice_cabecalho(linhas: list[tuple]) -> int:
    for i, linha in enumerate(linhas):
        if linha and linha[0] == "N":
            return i
    raise ValueError('Linha de cabeçalho ("N | Nome da Variável | ...") não encontrada.')


def _titulo(linhas: list[tuple]) -> str:
    for linha in linhas:
        for valor in linha or ():
            if valor and TITULO_PADRAO in str(valor):
                return espacos_normalizados(valor)
    return TITULO_PADRAO


def ler_aba(chave: str, planilha) -> TabelaDicionario:
    linhas = list(planilha.iter_rows(values_only=True))
    i_cabecalho = _indice_cabecalho(linhas)

    linha_anos = linhas[i_cabecalho + 1]
    colunas_ano: list[tuple[int, str]] = []
    j = _COLUNA_PRIMEIRO_ANO
    while j < len(linha_anos) and linha_anos[j] is not None:
        colunas_ano.append((j, f"20{str(linha_anos[j]).strip()}"))
        j += 1
    coluna_notas = j

    def celula(linha: tuple, i: int):
        return linha[i] if i < len(linha) else None

    variaveis = []
    for linha in linhas[i_cabecalho + 2:]:
        if not linha or not isinstance(linha[0], (int, float)):
            continue
        nome = str(linha[1] or "").strip()
        if not nome:
            continue
        rotulos, notas_aplicabilidade = parse_categoria(linha[5])
        notas = celula(linha, coluna_notas)
        variaveis.append(Variavel(
            ordem=int(linha[0]),
            nome=nome,
            descricao=espacos_normalizados(linha[2] or ""),
            tipo=str(linha[3]).strip() if linha[3] else "Num",
            tamanho=linha[4] if isinstance(linha[4], (int, float)) else None,
            rotulos_valor=rotulos,
            notas_aplicabilidade=notas_aplicabilidade,
            anos_coleta={
                ano: (str(v).strip().lower() if (v := celula(linha, i)) is not None else None)
                for i, ano in colunas_ano
            },
            notas=(espacos_normalizados(notas) or None) if notas else None,
        ))

    return TabelaDicionario(
        chave=chave,
        titulo=_titulo(linhas[:i_cabecalho]),
        variaveis=variaveis,
        ano_dicionario=colunas_ano[-1][1] if colunas_ano else None,
    )


def ler_dicionario(caminho: Path) -> dict[str, TabelaDicionario]:
    """Todas as abas reconhecidas, na ordem canônica das tabelas.

    Abas não reconhecidas são avisadas e ignoradas.
    """
    import openpyxl

    wb = openpyxl.load_workbook(caminho, data_only=True)
    abas: dict[str, str] = {}
    for nome_aba in wb.sheetnames:
        if tabela := identificar_tabela(nome_aba.strip()):
            abas.setdefault(tabela, nome_aba)
        else:
            log.warning('Aba "%s" não reconhecida como tabela do Censo; ignorando.', nome_aba)
    return {chave: ler_aba(chave, wb[abas[chave]]) for chave in CHAVES if chave in abas}


def tabelas_do_dicionario(fonte) -> list[str]:
    """Tabelas com aba no dicionário, em ordem canônica, lendo só os nomes das abas.

    `fonte`: caminho, `bytes` ou file-like. Arquivo ilegível → lista vazia.
    """
    import openpyxl

    if isinstance(fonte, (bytes, bytearray, memoryview)):
        fonte = io.BytesIO(bytes(fonte))
    try:
        wb = openpyxl.load_workbook(fonte, read_only=True)
    except Exception:
        return []
    try:
        achadas = {t for aba in wb.sheetnames if (t := identificar_tabela(aba.strip()))}
    finally:
        wb.close()
    return [t for t in CHAVES if t in achadas]
